"""Bounded, first-sheet XLSX reader for local commercial uploads.

Never execute formulas or macros. Only first worksheet and cached formula
values are considered. Reject compressed workbook expansion beyond budget.
"""
from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile, ZipFile

import pandas as pd

from backend.app.services.basic_report import MAX_COLUMNS, MAX_ROWS, MAX_UPLOAD_BYTES, validate_basic_frame


MAX_XLSX_INFLATED = 80 * 1024 * 1024


def read_xlsx_bounded(data: bytes) -> pd.DataFrame:
    if not data or len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("XLSX file is empty or exceeds 30 MiB.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 250 or not infos:
                raise ValueError("XLSX archive contains too many entries.")
            if sum(item.file_size for item in infos) > MAX_XLSX_INFLATED:
                raise ValueError("XLSX uncompressed data exceeds 80 MiB.")
            if any(item.file_size > 50 * 1024 * 1024 for item in infos):
                raise ValueError("XLSX contains an oversized part.")
            names_list = [item.filename.replace("\\", "/") for item in infos]
            names = set(names_list)
            lower_names = {name.lower() for name in names}
            if len(names) != len(names_list) or len(lower_names) != len(names):
                raise ValueError("Duplicate or case-colliding XLSX archive entries.")
            if "[content_types].xml" not in lower_names or "xl/workbook.xml" not in lower_names:
                raise ValueError("Invalid XLSX workbook structure.")
            if any("vbaproject" in name or name.endswith(".bin") for name in lower_names):
                raise ValueError("Macro or binary workbook parts are not supported.")
            if any(name.startswith("xl/externallinks/") for name in lower_names):
                raise ValueError("External-link XLSX workbooks are not supported.")
    except BadZipFile as exc:
        raise ValueError("Invalid XLSX archive.") from exc

    from openpyxl import load_workbook
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, ValueError, TypeError, KeyError) as exc:
        raise ValueError("Failed to open the XLSX workbook.") from exc

    try:
        if not workbook.worksheets:
            raise ValueError("XLSX has no worksheets.")
        sheet = workbook.worksheets[0]
        if sheet.max_column > MAX_COLUMNS or sheet.max_row > MAX_ROWS + 1:
            raise ValueError(f"XLSX limit exceeded: {MAX_ROWS} data rows and {MAX_COLUMNS} columns.")
        # Read actual dimensions, not the maximum budget width: openpyxl
        # pads rows to max_col and would mistake the padding for blank headers.
        records = sheet.iter_rows(values_only=True, max_row=sheet.max_row,
                                 max_col=sheet.max_column)
        first = next(records, None)
        if first is None or not first:
            raise ValueError("XLSX has no header row.")
        headers = [value for value in first if value is not None]
        if len(headers) != len(first):
            raise ValueError("XLSX header must not contain blank cells.")
        if not all(isinstance(value, str) and value.strip() for value in headers):
            raise ValueError("XLSX headers must be nonempty text.")
        if len(set(headers)) != len(headers):
            raise ValueError("Duplicate column names are not supported.")
        rows = []
        for row in records:
            if len(rows) >= MAX_ROWS:
                raise ValueError(f"XLSX limit exceeded: {MAX_ROWS} data rows.")
            if len(row) > len(headers) and any(value is not None for value in row[len(headers):]):
                raise ValueError("XLSX has cells outside the header width.")
            rows.append(tuple(row[:len(headers)]))
        df = pd.DataFrame(rows, columns=list(headers))
        df = df.dropna(how="all")
        validate_basic_frame(df)
        return df
    finally:
        workbook.close()


def read_uploaded_table(data: bytes, filename: str) -> pd.DataFrame:
    name = filename.lower()
    if name.endswith(".xlsx"):
        return read_xlsx_bounded(data)
    if name.endswith(".csv"):
        # Import lazily because source module also imports our parser.
        from backend.app.main import _read_csv_from_bytes
        frame = _read_csv_from_bytes(data)
        validate_basic_frame(frame)
        return frame
    raise ValueError("Only CSV and XLSX files are supported.")
