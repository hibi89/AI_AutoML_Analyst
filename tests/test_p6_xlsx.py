"""P6 XLSX ingestion: local bounded first-sheet workbook, no formula execution."""
from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pandas as pd
import pytest

from backend.app.main import app
from backend.app.services.tabular_input import read_xlsx_bounded


def make_xlsx(rows: int = 60, invalid_header: bool = False) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["visits", "revenue", None if invalid_header else "target"])
    for i in range(rows):
        sheet.append([i % 11, float(100 + i % 17), i % 2])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_local_basic_xlsx_produces_real_zip():
    data = make_xlsx()
    response = TestClient(app).post(
        "/api/basic-report",
        files={"file": ("data.xlsx", BytesIO(data),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    with ZipFile(BytesIO(response.content)) as z:
        assert z.read("report.pdf").startswith(b"%PDF-")
        import json
        item = json.loads(z.read("summary.json"))
        assert item["rows"] == 60
        assert item["columns"] == 3


def test_local_standard_xlsx_runs_holdout():
    response = TestClient(app).post(
        "/api/standard-report",
        data={"target_column": "target", "task_type": "classification"},
        files={"file": ("data.xlsx", BytesIO(make_xlsx()),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    with ZipFile(BytesIO(response.content)) as z:
        assert "model_scores.csv" in z.namelist()


def test_formula_cells_are_not_executed():
    book = Workbook()
    sheet = book.active
    sheet.append(["input", "target"])
    for i in range(6):
        sheet.append([i, "=SUM(1,2)" if i == 0 else float(i)])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    data = read_xlsx_bounded(buffer.getvalue())
    assert pd.isna(data.loc[0, "target"])
    assert data.loc[1, "target"] == 1.0


def test_invalid_workbook_headers_and_archive_rejected():
    with pytest.raises(ValueError, match="header"):
        read_xlsx_bounded(make_xlsx(invalid_header=True))
    with pytest.raises(ValueError, match="Invalid XLSX"):
        read_xlsx_bounded(b"hello")
    out = BytesIO()
    with ZipFile(out, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "ok")
        archive.writestr("xl/workbook.xml", "ok")
        archive.writestr("xl/vbaProject.bin", b"unsafe macro")
    with pytest.raises(ValueError, match="Macro"):
        read_xlsx_bounded(out.getvalue())


def test_xlsx_column_limit_is_bounded():
    book = Workbook()
    sheet = book.active
    sheet.append([f"Col{i}" for i in range(51)])
    sheet.append([i for i in range(51)])
    output = BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(ValueError, match="limit"):
        read_xlsx_bounded(output.getvalue())


def test_no_external_host_can_upload_xlsx():
    response = TestClient(app).post(
        "/api/basic-report",
        headers={"host": "evil.example"},
        files={"file": ("data.xlsx", BytesIO(make_xlsx()), "application/octet-stream")},
    )
    assert response.status_code == 403
